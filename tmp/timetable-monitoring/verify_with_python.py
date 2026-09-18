from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

source_path = Path(r"D:\Downloads\MakaLearn Proposed Timetable Monitoring (1).xlsx")
output_path = Path(r"D:\MakaLearn\outputs\timetable-commit-review-20260918\MakaLearn Proposed Timetable Monitoring - Commit Completion.xlsx")

source = load_workbook(source_path, data_only=False, read_only=False)
output = load_workbook(output_path, data_only=False, read_only=False)

assert source.sheetnames == output.sheetnames, "Sheet names or order changed"
assert len(output.sheetnames) == 18, "Unexpected worksheet count"

source_sheet = source["Progress Table"]
output_sheet = output["Progress Table"]

for row in range(1, 154):
    for column in range(1, 6):
        assert source_sheet.cell(row, column).value == output_sheet.cell(row, column).value, (
            f"Source timetable value changed at row {row}, column {column}"
        )

statuses = [output_sheet.cell(row, 7).value for row in range(2, 137)]
dates = [output_sheet.cell(row, 6).value for row in range(2, 137)]
comments = [output_sheet.cell(row, 8).value for row in range(2, 137)]
counts = Counter(statuses)

assert all(status in {"Finished", "Partially finished", "Not yet evidenced"} for status in statuses)
assert all(comments), "Missing evidence comment"
assert sum(isinstance(value, datetime) for value in dates) == counts["Finished"]
assert all(
    output_sheet.cell(row, 6).number_format == "mmmm d, yyyy"
    for row in range(2, 137)
), "Completion-date format is inconsistent"

formula_errors = []
for row in output_sheet.iter_rows(min_row=1, max_row=153, min_col=1, max_col=8):
    for cell in row:
        if isinstance(cell.value, str) and cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")):
            formula_errors.append(cell.coordinate)

print(f"sheets={len(output.sheetnames)}")
print(f"counts={dict(counts)}")
print(f"dated_rows={sum(isinstance(value, datetime) for value in dates)}")
print(f"progress_table_formula_errors={formula_errors}")
print(f"source_bytes={source_path.stat().st_size}")
print(f"output_bytes={output_path.stat().st_size}")
print(f"verified={output_path}")
